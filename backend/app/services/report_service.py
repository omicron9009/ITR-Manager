"""Service — Comprehensive report builder for DASHBOARD_USER / Partner."""

from datetime import datetime, timezone
from io import BytesIO
from typing import Optional
from uuid import UUID

from sqlalchemy import extract, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.enums import FilingStatus, TagType, UserRole
from app.models.executive_tag import ExecutiveTag
from app.models.filing import ITRFiling
from app.models.manager_client_assignment import ManagerClientAssignment
from app.models.manager_executive_assignment import ManagerExecutiveAssignment
from app.models.tag import Tag
from app.models.user import User


# ═══════════════════════════════════════════════════════════════
# HELPERS
# ═══════════════════════════════════════════════════════════════


async def _avg_completion_days(db: AsyncSession, filters: list) -> Optional[float]:
    """Calculate average days from initiated_at to completed_at with optional filters."""
    result = await db.execute(
        select(
            func.avg(extract("epoch", ITRFiling.completed_at - ITRFiling.initiated_at) / 86400)
        ).where(ITRFiling.completed_at.isnot(None), *filters)
    )
    val = result.scalar()
    return round(float(val), 1) if val else None


async def _filing_stats_for_ids(db: AsyncSession, executive_ids: list[UUID], fy: Optional[str] = None) -> dict:
    """Aggregate filing stats for a set of executive IDs."""
    filters = [ITRFiling.assigned_executive_id.in_(executive_ids)]
    if fy:
        filters.append(ITRFiling.financial_year == fy)

    result = await db.execute(
        select(ITRFiling.status, func.count(ITRFiling.id))
        .where(*filters)
        .group_by(ITRFiling.status)
    )
    status_map = {row[0]: row[1] for row in result.all()}

    total = sum(status_map.values())
    completed = status_map.get(FilingStatus.COMPLETED, 0)
    halted = status_map.get(FilingStatus.HALTED, 0)
    active = total - completed - halted

    return {"total": total, "active": active, "completed": completed, "halted": halted}


# ═══════════════════════════════════════════════════════════════
# MAIN BUILDER
# ═══════════════════════════════════════════════════════════════


async def build_comprehensive_report(db: AsyncSession, financial_year: Optional[str] = None) -> dict:
    """Build full comprehensive report payload."""
    now = datetime.now(timezone.utc)

    # ── Available FYs ──
    fy_result = await db.execute(
        select(ITRFiling.financial_year).distinct().order_by(ITRFiling.financial_year.desc())
    )
    all_fys = [row[0] for row in fy_result.all()]

    # ── Base filter for optional FY ──
    fy_filter = [ITRFiling.financial_year == financial_year] if financial_year else []

    # ═══ OVERALL SUMMARY ═══
    overall = await _build_overall(db, fy_filter)

    # ═══ FY-WISE ═══
    fy_wise = await _build_fy_wise(db, financial_year)

    # ═══ MANAGER DISTRIBUTION ═══
    manager_dist = await _build_manager_distribution(db, financial_year)

    # ═══ LOCATION DISTRIBUTION ═══
    location_dist = await _build_location_distribution(db, financial_year)

    # ═══ PENDING REPORT ═══
    pending = await _build_pending_report(db, fy_filter)

    # ═══ LEADERBOARDS ═══
    lb_executive = await _build_executive_leaderboard(db, fy_filter)
    lb_manager = await _build_manager_leaderboard(db, financial_year)
    lb_location = await _build_location_leaderboard(db, financial_year)

    return {
        "generated_at": now,
        "financial_years": all_fys,
        "filtered_fy": financial_year,
        "overall": overall,
        "fy_wise": fy_wise,
        "manager_distribution": manager_dist,
        "location_distribution": location_dist,
        "pending_report": pending,
        "leaderboard_executive": lb_executive,
        "leaderboard_manager": lb_manager,
        "leaderboard_location": lb_location,
    }


# ═══════════════════════════════════════════════════════════════
# SECTION BUILDERS
# ═══════════════════════════════════════════════════════════════


async def _build_overall(db: AsyncSession, fy_filter: list) -> dict:
    """Build overall summary section."""
    # Client counts
    client_result = await db.execute(
        select(
            func.count(User.id).label("total"),
            func.count(User.id).filter(User.account_status == "ACTIVE").label("active"),
        ).where(User.role == UserRole.CLIENT)
    )
    cr = client_result.one()

    # Executive counts
    exec_result = await db.execute(
        select(
            func.count(User.id).label("total"),
            func.count(User.id).filter(User.is_active == True).label("active"),
        ).where(User.role == UserRole.EXECUTIVE)
    )
    er = exec_result.one()

    # Filing status breakdown
    status_result = await db.execute(
        select(ITRFiling.status, func.count(ITRFiling.id))
        .where(*fy_filter)
        .group_by(ITRFiling.status)
    )
    status_map = {row[0].value: row[1] for row in status_result.all()}
    total_filings = sum(status_map.values())
    completed = status_map.get("COMPLETED", 0)
    halted = status_map.get("HALTED", 0)
    active = total_filings - completed - halted

    breakdown = [{"status": s, "count": c} for s, c in status_map.items()]

    # Avg completion
    avg_total = await _avg_completion_days(db, fy_filter)

    # Stage timings
    avg_doc = await db.execute(
        select(func.avg(extract("epoch", ITRFiling.documents_approved_at - ITRFiling.documents_submitted_at) / 86400))
        .where(ITRFiling.documents_approved_at.isnot(None), ITRFiling.documents_submitted_at.isnot(None), *fy_filter)
    )
    avg_comp = await db.execute(
        select(func.avg(extract("epoch", ITRFiling.computation_approved_at - ITRFiling.computation_uploaded_at) / 86400))
        .where(ITRFiling.computation_approved_at.isnot(None), ITRFiling.computation_uploaded_at.isnot(None), *fy_filter)
    )
    avg_filing = await db.execute(
        select(func.avg(extract("epoch", ITRFiling.filed_at - ITRFiling.computation_approved_at) / 86400))
        .where(ITRFiling.filed_at.isnot(None), ITRFiling.computation_approved_at.isnot(None), *fy_filter)
    )
    avg_payment = await db.execute(
        select(func.avg(extract("epoch", ITRFiling.payment_received_at - ITRFiling.filed_at) / 86400))
        .where(ITRFiling.payment_received_at.isnot(None), ITRFiling.filed_at.isnot(None), *fy_filter)
    )

    def _round_scalar(r):
        v = r.scalar()
        return round(float(v), 1) if v else None

    return {
        "total_clients": cr.total,
        "active_clients": cr.active,
        "total_executives": er.total,
        "active_executives": er.active,
        "total_filings": total_filings,
        "active_filings": active,
        "completed_filings": completed,
        "halted_filings": halted,
        "filing_status_breakdown": breakdown,
        "avg_days_to_complete": avg_total,
        "avg_days_by_stage": {
            "avg_days_document_processing": _round_scalar(avg_doc),
            "avg_days_computation": _round_scalar(avg_comp),
            "avg_days_filing": _round_scalar(avg_filing),
            "avg_days_payment": _round_scalar(avg_payment),
        },
    }


async def _build_fy_wise(db: AsyncSession, financial_year: Optional[str]) -> list[dict]:
    """Build FY-wise summary with client drill-down."""
    # Determine which FYs to include
    if financial_year:
        target_fys = [financial_year]
    else:
        fy_result = await db.execute(
            select(ITRFiling.financial_year).distinct().order_by(ITRFiling.financial_year.desc())
        )
        target_fys = [row[0] for row in fy_result.all()]

    fy_items = []
    for fy in target_fys:
        # Status breakdown
        status_result = await db.execute(
            select(ITRFiling.status, func.count(ITRFiling.id))
            .where(ITRFiling.financial_year == fy)
            .group_by(ITRFiling.status)
        )
        status_map = {row[0].value: row[1] for row in status_result.all()}
        total = sum(status_map.values())
        completed = status_map.get("COMPLETED", 0)
        halted = status_map.get("HALTED", 0)
        active = total - completed - halted

        # Avg completion for this FY
        avg_days = await _avg_completion_days(db, [ITRFiling.financial_year == fy])

        # Client drill-down (all filings for this FY)
        filings_result = await db.execute(
            select(ITRFiling).where(ITRFiling.financial_year == fy)
            .order_by(ITRFiling.updated_at.desc())
        )
        filings = filings_result.scalars().all()

        clients = []
        for f in filings:
            c_result = await db.execute(select(User).where(User.id == f.client_id))
            client = c_result.scalar_one_or_none()
            exec_name = None
            if f.assigned_executive_id:
                e_result = await db.execute(select(User.full_name).where(User.id == f.assigned_executive_id))
                exec_name = e_result.scalar()
            clients.append({
                "filing_id": f.id,
                "client_id": f.client_id,
                "client_name": client.full_name if client else "Unknown",
                "client_email": client.email if client else "",
                "financial_year": f.financial_year,
                "status": f.status.value,
                "assigned_executive_name": exec_name,
                "initiated_at": f.initiated_at,
                "last_updated": f.updated_at,
            })

        fy_items.append({
            "financial_year": fy,
            "total_filings": total,
            "active_filings": active,
            "completed_filings": completed,
            "halted_filings": halted,
            "filing_status_breakdown": [{"status": s, "count": c} for s, c in status_map.items()],
            "avg_days_to_complete": avg_days,
            "clients": clients,
        })

    return fy_items


async def _build_manager_distribution(db: AsyncSession, financial_year: Optional[str]) -> list[dict]:
    """Build manager-wise distribution using real Manager role assignments."""
    mgr_result = await db.execute(
        select(User).where(User.role == UserRole.MANAGER, User.is_active == True).order_by(User.full_name)
    )
    managers = mgr_result.scalars().all()

    fy_filter = [ITRFiling.financial_year == financial_year] if financial_year else []
    items = []

    for mgr in managers:
        # Get executive IDs for this manager
        et_result = await db.execute(
            select(ManagerExecutiveAssignment.executive_id).where(
                ManagerExecutiveAssignment.manager_id == mgr.id,
                ManagerExecutiveAssignment.is_active == True,
            )
        )
        exec_ids = [row[0] for row in et_result.all()]

        if not exec_ids:
            items.append({
                "manager_id": mgr.id,
                "manager_name": mgr.full_name,
                "executive_count": 0,
                "total_filings": 0,
                "active_filings": 0,
                "completed_filings": 0,
                "halted_filings": 0,
                "avg_days_to_complete": None,
                "executives": [],
            })
            continue

        # Aggregate stats
        stats = await _filing_stats_for_ids(db, exec_ids, financial_year)
        avg_days = await _avg_completion_days(
            db, [ITRFiling.assigned_executive_id.in_(exec_ids)] + fy_filter
        )

        # Per-executive breakdown
        executives = []
        for eid in exec_ids:
            user_result = await db.execute(select(User).where(User.id == eid))
            user = user_result.scalar_one_or_none()
            if not user:
                continue
            e_stats = await _filing_stats_for_ids(db, [eid], financial_year)
            e_avg = await _avg_completion_days(
                db, [ITRFiling.assigned_executive_id == eid] + fy_filter
            )
            executives.append({
                "executive_id": user.id,
                "executive_name": user.full_name,
                "total_filings": e_stats["total"],
                "active_filings": e_stats["active"],
                "completed_filings": e_stats["completed"],
                "halted_filings": e_stats["halted"],
                "avg_days_to_complete": e_avg,
            })

        items.append({
            "manager_id": mgr.id,
            "manager_name": mgr.full_name,
            "executive_count": len(exec_ids),
            "total_filings": stats["total"],
            "active_filings": stats["active"],
            "completed_filings": stats["completed"],
            "halted_filings": stats["halted"],
            "avg_days_to_complete": avg_days,
            "executives": executives,
        })

    return items


async def _build_location_distribution(db: AsyncSession, financial_year: Optional[str]) -> list[dict]:
    """Build location distribution with executives breakdown."""
    loc_result = await db.execute(
        select(Tag).where(Tag.tag_type == TagType.LOCATION, Tag.is_active == True).order_by(Tag.name)
    )
    location_tags = loc_result.scalars().all()

    fy_filter = [ITRFiling.financial_year == financial_year] if financial_year else []
    items = []

    for loc_tag in location_tags:
        # Executives in this location
        et_result = await db.execute(
            select(ExecutiveTag.executive_id).where(
                ExecutiveTag.tag_id == loc_tag.id, ExecutiveTag.is_active == True
            )
        )
        loc_exec_ids = set(row[0] for row in et_result.all())

        if not loc_exec_ids:
            items.append({
                "tag_id": loc_tag.id,
                "location_name": loc_tag.name,
                "executive_count": 0,
                "manager_count": 0,
                "total_filings": 0,
                "active_filings": 0,
                "completed_filings": 0,
                "halted_filings": 0,
                "avg_days_to_complete": None,
                "executives": [],
            })
            continue

        # Stats for the whole location
        stats = await _filing_stats_for_ids(db, list(loc_exec_ids), financial_year)
        avg_days = await _avg_completion_days(
            db, [ITRFiling.assigned_executive_id.in_(loc_exec_ids)] + fy_filter
        )

        # Count distinct managers assigned to executives in this location
        mgr_result = await db.execute(
            select(ManagerExecutiveAssignment.manager_id).where(
                ManagerExecutiveAssignment.executive_id.in_(loc_exec_ids),
                ManagerExecutiveAssignment.is_active == True,
            ).distinct()
        )
        manager_count = len(mgr_result.scalars().all())

        # Per-executive breakdown
        executives = []
        for eid in loc_exec_ids:
            user_result = await db.execute(select(User).where(User.id == eid))
            user = user_result.scalar_one_or_none()
            if not user:
                continue
            e_stats = await _filing_stats_for_ids(db, [eid], financial_year)
            e_avg = await _avg_completion_days(
                db, [ITRFiling.assigned_executive_id == eid] + fy_filter
            )
            executives.append({
                "executive_id": user.id,
                "executive_name": user.full_name,
                "total_filings": e_stats["total"],
                "active_filings": e_stats["active"],
                "completed_filings": e_stats["completed"],
                "halted_filings": e_stats["halted"],
                "avg_days_to_complete": e_avg,
            })

        items.append({
            "tag_id": loc_tag.id,
            "location_name": loc_tag.name,
            "executive_count": len(loc_exec_ids),
            "manager_count": manager_count,
            "total_filings": stats["total"],
            "active_filings": stats["active"],
            "completed_filings": stats["completed"],
            "halted_filings": stats["halted"],
            "avg_days_to_complete": avg_days,
            "executives": executives,
        })

    return items


async def _build_pending_report(db: AsyncSession, fy_filter: list) -> list[dict]:
    """Build pending filings report sorted by longest-pending first."""
    now = datetime.now(timezone.utc)

    # Get all non-completed, non-halted filings
    filings_result = await db.execute(
        select(ITRFiling).where(
            ITRFiling.status.notin_([FilingStatus.COMPLETED, FilingStatus.HALTED]),
            *fy_filter,
        ).order_by(ITRFiling.initiated_at.asc())  # oldest first = longest pending
    )
    filings = filings_result.scalars().all()

    items = []
    for f in filings:
        # Client info
        client_result = await db.execute(select(User).where(User.id == f.client_id))
        client = client_result.scalar_one_or_none()

        # Executive name
        exec_name = None
        if f.assigned_executive_id:
            e_result = await db.execute(select(User.full_name).where(User.id == f.assigned_executive_id))
            exec_name = e_result.scalar()

        # Tags for assigned executive
        location_tag_name = None
        if f.assigned_executive_id:
            tag_result = await db.execute(
                select(Tag.name, Tag.tag_type).join(ExecutiveTag, ExecutiveTag.tag_id == Tag.id).where(
                    ExecutiveTag.executive_id == f.assigned_executive_id,
                    ExecutiveTag.is_active == True,
                    Tag.is_active == True,
                )
            )
            for tag_name, tag_type in tag_result.all():
                if tag_type == TagType.LOCATION and not location_tag_name:
                    location_tag_name = tag_name

        # Get manager name from real assignment
        manager_name = None
        if f.assigned_executive_id:
            mgr_assign_result = await db.execute(
                select(User.full_name).join(
                    ManagerExecutiveAssignment,
                    ManagerExecutiveAssignment.manager_id == User.id,
                ).where(
                    ManagerExecutiveAssignment.executive_id == f.assigned_executive_id,
                    ManagerExecutiveAssignment.is_active == True,
                )
            )
            mgr_row = mgr_assign_result.first()
            if mgr_row:
                manager_name = mgr_row[0]

        days_pending = (now - f.initiated_at.replace(tzinfo=timezone.utc)).total_seconds() / 86400

        items.append({
            "filing_id": f.id,
            "client_id": f.client_id,
            "client_name": client.full_name if client else "Unknown",
            "client_email": client.email if client else "",
            "financial_year": f.financial_year,
            "status": f.status.value,
            "assigned_executive_name": exec_name,
            "manager_name": manager_name,
            "location_tag": location_tag_name,
            "days_pending": round(days_pending, 1),
            "initiated_at": f.initiated_at,
            "last_updated": f.updated_at,
        })

    return items


async def _build_executive_leaderboard(db: AsyncSession, fy_filter: list) -> list[dict]:
    """Build executive leaderboard ranked by completed filings."""
    exec_result = await db.execute(
        select(User).where(User.role == UserRole.EXECUTIVE, User.is_active == True).order_by(User.full_name)
    )
    executives = exec_result.scalars().all()

    leaderboard = []
    for ex in executives:
        # Filing counts
        stats_result = await db.execute(
            select(ITRFiling.status, func.count(ITRFiling.id))
            .where(ITRFiling.assigned_executive_id == ex.id, *fy_filter)
            .group_by(ITRFiling.status)
        )
        status_map = {row[0]: row[1] for row in stats_result.all()}
        total = sum(status_map.values())
        completed = status_map.get(FilingStatus.COMPLETED, 0)
        halted = status_map.get(FilingStatus.HALTED, 0)
        active = total - completed - halted

        avg_days = await _avg_completion_days(
            db, [ITRFiling.assigned_executive_id == ex.id] + fy_filter
        )

        leaderboard.append({
            "executive_id": ex.id,
            "executive_name": ex.full_name,
            "completed_filings": completed,
            "active_filings": active,
            "total_filings": total,
            "avg_days_to_complete": avg_days,
        })

    # Sort by completed filings descending, then by avg_days ascending
    leaderboard.sort(key=lambda x: (-x["completed_filings"], x["avg_days_to_complete"] or 9999))

    # Assign ranks
    for i, item in enumerate(leaderboard, 1):
        item["rank"] = i

    return leaderboard


async def _build_manager_leaderboard(db: AsyncSession, financial_year: Optional[str]) -> list[dict]:
    """Build manager leaderboard using real Manager role assignments."""
    mgr_result = await db.execute(
        select(User).where(User.role == UserRole.MANAGER, User.is_active == True)
    )
    managers = mgr_result.scalars().all()

    fy_filter = [ITRFiling.financial_year == financial_year] if financial_year else []
    leaderboard = []

    for mgr in managers:
        et_result = await db.execute(
            select(ManagerExecutiveAssignment.executive_id).where(
                ManagerExecutiveAssignment.manager_id == mgr.id,
                ManagerExecutiveAssignment.is_active == True,
            )
        )
        exec_ids = [row[0] for row in et_result.all()]

        if not exec_ids:
            leaderboard.append({
                "manager_id": mgr.id,
                "manager_name": mgr.full_name,
                "executive_count": 0,
                "completed_filings": 0,
                "total_filings": 0,
                "avg_days_to_complete": None,
            })
            continue

        stats = await _filing_stats_for_ids(db, exec_ids, financial_year)
        avg_days = await _avg_completion_days(
            db, [ITRFiling.assigned_executive_id.in_(exec_ids)] + fy_filter
        )

        leaderboard.append({
            "manager_id": mgr.id,
            "manager_name": mgr.full_name,
            "executive_count": len(exec_ids),
            "completed_filings": stats["completed"],
            "total_filings": stats["total"],
            "avg_days_to_complete": avg_days,
        })

    leaderboard.sort(key=lambda x: (-x["completed_filings"], x["avg_days_to_complete"] or 9999))
    for i, item in enumerate(leaderboard, 1):
        item["rank"] = i

    return leaderboard


async def _build_location_leaderboard(db: AsyncSession, financial_year: Optional[str]) -> list[dict]:
    """Build location leaderboard ranked by completed filings."""
    tags_result = await db.execute(
        select(Tag).where(Tag.tag_type == TagType.LOCATION, Tag.is_active == True)
    )
    location_tags = tags_result.scalars().all()

    fy_filter = [ITRFiling.financial_year == financial_year] if financial_year else []
    leaderboard = []

    for tag in location_tags:
        et_result = await db.execute(
            select(ExecutiveTag.executive_id).where(
                ExecutiveTag.tag_id == tag.id, ExecutiveTag.is_active == True
            )
        )
        exec_ids = [row[0] for row in et_result.all()]

        if not exec_ids:
            leaderboard.append({
                "tag_id": tag.id,
                "location_name": tag.name,
                "executive_count": 0,
                "completed_filings": 0,
                "total_filings": 0,
                "avg_days_to_complete": None,
            })
            continue

        stats = await _filing_stats_for_ids(db, exec_ids, financial_year)
        avg_days = await _avg_completion_days(
            db, [ITRFiling.assigned_executive_id.in_(exec_ids)] + fy_filter
        )

        leaderboard.append({
            "tag_id": tag.id,
            "location_name": tag.name,
            "executive_count": len(exec_ids),
            "completed_filings": stats["completed"],
            "total_filings": stats["total"],
            "avg_days_to_complete": avg_days,
        })

    leaderboard.sort(key=lambda x: (-x["completed_filings"], x["avg_days_to_complete"] or 9999))
    for i, item in enumerate(leaderboard, 1):
        item["rank"] = i

    return leaderboard


# ═══════════════════════════════════════════════════════════════
# EXCEL REPORT GENERATOR
# ═══════════════════════════════════════════════════════════════


async def generate_excel_report(db: AsyncSession, financial_year: Optional[str] = None) -> BytesIO:
    """Generate a multi-sheet Excel workbook from the comprehensive report data."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    data = await build_comprehensive_report(db, financial_year)
    wb = Workbook()

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    def _write_headers(ws, headers: list[str]):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

    # ── Sheet 1: Overall Summary ──
    ws = wb.active
    ws.title = "Overall"
    overall = data["overall"]
    summary_rows = [
        ("Total Clients", overall["total_clients"]),
        ("Active Clients", overall["active_clients"]),
        ("Total Executives", overall["total_executives"]),
        ("Active Executives", overall["active_executives"]),
        ("Total Filings", overall["total_filings"]),
        ("Active Filings", overall["active_filings"]),
        ("Completed Filings", overall["completed_filings"]),
        ("Halted Filings", overall["halted_filings"]),
        ("Avg Days to Complete", overall["avg_days_to_complete"] or "N/A"),
        ("Avg Days Document Processing", overall["avg_days_by_stage"].get("avg_days_document_processing") or "N/A"),
        ("Avg Days Computation", overall["avg_days_by_stage"].get("avg_days_computation") or "N/A"),
        ("Avg Days Filing", overall["avg_days_by_stage"].get("avg_days_filing") or "N/A"),
        ("Avg Days Payment", overall["avg_days_by_stage"].get("avg_days_payment") or "N/A"),
    ]
    _write_headers(ws, ["Metric", "Value"])
    for row_idx, (metric, value) in enumerate(summary_rows, 2):
        ws.cell(row=row_idx, column=1, value=metric)
        ws.cell(row=row_idx, column=2, value=value)

    # ── Sheet 2: FY-Wise ──
    ws2 = wb.create_sheet("FY-Wise")
    _write_headers(ws2, ["Financial Year", "Total", "Active", "Completed", "Halted", "Avg Days"])
    for row_idx, fy in enumerate(data["fy_wise"], 2):
        ws2.cell(row=row_idx, column=1, value=fy["financial_year"])
        ws2.cell(row=row_idx, column=2, value=fy["total_filings"])
        ws2.cell(row=row_idx, column=3, value=fy["active_filings"])
        ws2.cell(row=row_idx, column=4, value=fy["completed_filings"])
        ws2.cell(row=row_idx, column=5, value=fy["halted_filings"])
        ws2.cell(row=row_idx, column=6, value=fy["avg_days_to_complete"] or "N/A")

    # ── Sheet 3: Manager Distribution ──
    ws3 = wb.create_sheet("Manager Distribution")
    _write_headers(ws3, ["Manager", "Executives", "Total Filings", "Active", "Completed", "Halted", "Avg Days"])
    row_idx = 2
    for mgr in data["manager_distribution"]:
        ws3.cell(row=row_idx, column=1, value=mgr["manager_name"])
        ws3.cell(row=row_idx, column=2, value=mgr["executive_count"])
        ws3.cell(row=row_idx, column=3, value=mgr["total_filings"])
        ws3.cell(row=row_idx, column=4, value=mgr["active_filings"])
        ws3.cell(row=row_idx, column=5, value=mgr["completed_filings"])
        ws3.cell(row=row_idx, column=6, value=mgr["halted_filings"])
        ws3.cell(row=row_idx, column=7, value=mgr["avg_days_to_complete"] or "N/A")
        row_idx += 1

    # ── Sheet 4: Location Distribution ──
    ws4 = wb.create_sheet("Location Distribution")
    _write_headers(ws4, ["Location", "Executives", "Managers", "Total Filings", "Active", "Completed", "Halted", "Avg Days"])
    row_idx = 2
    for loc in data["location_distribution"]:
        ws4.cell(row=row_idx, column=1, value=loc["location_name"])
        ws4.cell(row=row_idx, column=2, value=loc["executive_count"])
        ws4.cell(row=row_idx, column=3, value=loc["manager_count"])
        ws4.cell(row=row_idx, column=4, value=loc["total_filings"])
        ws4.cell(row=row_idx, column=5, value=loc["active_filings"])
        ws4.cell(row=row_idx, column=6, value=loc["completed_filings"])
        ws4.cell(row=row_idx, column=7, value=loc["halted_filings"])
        ws4.cell(row=row_idx, column=8, value=loc["avg_days_to_complete"] or "N/A")
        row_idx += 1

    # ── Sheet 5: Pending Report ──
    ws5 = wb.create_sheet("Pending Report")
    _write_headers(ws5, ["Client Name", "Client Email", "FY", "Status", "Executive", "Manager", "Location", "Days Pending", "Initiated"])
    row_idx = 2
    for p in data["pending_report"]:
        ws5.cell(row=row_idx, column=1, value=p["client_name"])
        ws5.cell(row=row_idx, column=2, value=p["client_email"])
        ws5.cell(row=row_idx, column=3, value=p["financial_year"])
        ws5.cell(row=row_idx, column=4, value=p["status"])
        ws5.cell(row=row_idx, column=5, value=p["assigned_executive_name"] or "Unassigned")
        ws5.cell(row=row_idx, column=6, value=p["manager_name"] or "N/A")
        ws5.cell(row=row_idx, column=7, value=p["location_tag"] or "N/A")
        ws5.cell(row=row_idx, column=8, value=p["days_pending"])
        ws5.cell(row=row_idx, column=9, value=p["initiated_at"].strftime("%Y-%m-%d") if p["initiated_at"] else "")
        row_idx += 1

    # ── Sheet 6: Executive Leaderboard ──
    ws6 = wb.create_sheet("Executive Leaderboard")
    _write_headers(ws6, ["Rank", "Executive Name", "Completed", "Active", "Total", "Avg Days"])
    for row_idx, item in enumerate(data["leaderboard_executive"], 2):
        ws6.cell(row=row_idx, column=1, value=item["rank"])
        ws6.cell(row=row_idx, column=2, value=item["executive_name"])
        ws6.cell(row=row_idx, column=3, value=item["completed_filings"])
        ws6.cell(row=row_idx, column=4, value=item["active_filings"])
        ws6.cell(row=row_idx, column=5, value=item["total_filings"])
        ws6.cell(row=row_idx, column=6, value=item["avg_days_to_complete"] or "N/A")

    # ── Sheet 7: Manager Leaderboard ──
    ws7 = wb.create_sheet("Manager Leaderboard")
    _write_headers(ws7, ["Rank", "Manager", "Executives", "Completed", "Total", "Avg Days"])
    for row_idx, item in enumerate(data["leaderboard_manager"], 2):
        ws7.cell(row=row_idx, column=1, value=item["rank"])
        ws7.cell(row=row_idx, column=2, value=item["manager_name"])
        ws7.cell(row=row_idx, column=3, value=item["executive_count"])
        ws7.cell(row=row_idx, column=4, value=item["completed_filings"])
        ws7.cell(row=row_idx, column=5, value=item["total_filings"])
        ws7.cell(row=row_idx, column=6, value=item["avg_days_to_complete"] or "N/A")

    # ── Sheet 8: Location Leaderboard ──
    ws8 = wb.create_sheet("Location Leaderboard")
    _write_headers(ws8, ["Rank", "Location", "Executives", "Completed", "Total", "Avg Days"])
    for row_idx, item in enumerate(data["leaderboard_location"], 2):
        ws8.cell(row=row_idx, column=1, value=item["rank"])
        ws8.cell(row=row_idx, column=2, value=item["location_name"])
        ws8.cell(row=row_idx, column=3, value=item["executive_count"])
        ws8.cell(row=row_idx, column=4, value=item["completed_filings"])
        ws8.cell(row=row_idx, column=5, value=item["total_filings"])
        ws8.cell(row=row_idx, column=6, value=item["avg_days_to_complete"] or "N/A")

    # Auto-size columns (approximate)
    for ws_item in wb.worksheets:
        for col in ws_item.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws_item.column_dimensions[col_letter].width = min(max_length + 2, 40)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
